"""RAG indexer — load, clean, chunk, embed, and store documents in KB.

Pipeline stages:
  1. Load document (PDF/DOCX/HTML/TXT)
  2. Clean text (normalise whitespace, remove artifacts)
  3. Chunk text (recursive split with overlap)
  4. Embed chunks via RouterAI bge-m3
  5. Store chunks + embeddings in kb_chunks table
"""

from __future__ import annotations

import asyncio
import hashlib
import logging
import re
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from agent.attachment_parser import _parse_pdf, _parse_docx
from agent.llm_client import embed_batch
from app.config import settings
from app.models.kb_chunk import KBChunk
from app.models.knowledge_base import KnowledgeBaseArticle

logger = logging.getLogger(__name__)

# Thread pool for CPU-bound file parsing (PDF OCR, DOCX, Excel)
_file_parse_pool = ThreadPoolExecutor(max_workers=4, thread_name_prefix="indexer")
# Semaphore limits concurrent file-parse tasks to avoid memory spikes
_parse_semaphore = asyncio.Semaphore(4)


# ---------------------------------------------------------------------------
# Stage 1 — Loader
# ---------------------------------------------------------------------------

def load_document(filepath: str) -> str:
    """Load a document and return raw text."""
    path = Path(filepath)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _parse_pdf(filepath)
    elif suffix in (".docx", ".doc"):
        return _parse_docx(filepath)
    elif suffix in (".html", ".htm"):
        return _load_html(filepath)
    elif suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")
    elif suffix in (".xls", ".xlsx"):
        return _parse_excel(filepath)
    else:
        raise ValueError(f"Unsupported file format: {suffix}")


async def load_document_async(filepath: str) -> str:
    """Load a document in a thread pool to avoid blocking the event loop.

    Uses a semaphore to cap concurrent file-parse operations (prevents
    memory spikes from multiple simultaneous OCR / large-file loads).
    """
    loop = asyncio.get_running_loop()
    async with _parse_semaphore:
        return await loop.run_in_executor(_file_parse_pool, load_document, filepath)


def compute_content_hash(filepath: str) -> str:
    """Compute SHA-256 hash of a file for change detection."""
    h = hashlib.sha256()
    with open(filepath, "rb") as f:
        while True:
            chunk = f.read(8192)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def compute_text_hash(text: str) -> str:
    """Compute SHA-256 hash of text content for change detection."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _load_html(filepath: str) -> str:
    """Load HTML and strip tags."""
    from bs4 import BeautifulSoup

    raw = Path(filepath).read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return soup.get_text(separator="\n")


def _parse_excel(filepath: str) -> str:
    """Extract text from XLS/XLSX files using openpyxl (xlsx) or xlrd (xls)."""
    path = Path(filepath)
    rows_text: list[str] = []

    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(filepath, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows_text.append(f"--- Лист: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None]
                if cells:
                    rows_text.append(" | ".join(cells))
        wb.close()
    else:
        # .xls via xlrd
        try:
            import xlrd

            wb = xlrd.open_workbook(filepath)
            for sheet in wb.sheets():
                rows_text.append(f"--- Лист: {sheet.name} ---")
                for rx in range(sheet.nrows):
                    cells = [
                        str(sheet.cell_value(rx, cx)).strip()
                        for cx in range(sheet.ncols)
                        if sheet.cell_value(rx, cx)
                    ]
                    if cells:
                        rows_text.append(" | ".join(cells))
        except ImportError:
            logger.warning(
                "xlrd not installed — cannot parse .xls file: %s", filepath,
            )
            return ""

    return "\n".join(rows_text)


# ---------------------------------------------------------------------------
# Stage 2 — Cleaner
# ---------------------------------------------------------------------------

def clean_text(raw: str) -> str:
    """Normalise and clean extracted text."""
    # NFC normalisation
    text = unicodedata.normalize("NFC", raw)
    # Remove invisible chars
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    # Collapse whitespace lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Clean excess spaces
    text = re.sub(r"[ \t]+", " ", text)
    # Strip header/footer page numbers
    text = re.sub(r"(?m)^\s*-?\s*\d+\s*-?\s*$", "", text)
    return text.strip()


# ---------------------------------------------------------------------------
# Stage 3 — Chunker (recursive text splitting)
# ---------------------------------------------------------------------------

def chunk_text(
    text: str,
    chunk_size: int = 1200,
    overlap: int = 100,
) -> list[str]:
    """Split text into overlapping chunks using recursive strategy.

    Priority of split separators: paragraph → line → sentence → word.
    """
    if len(text) <= chunk_size:
        return [text] if text.strip() else []

    separators = ["\n\n", "\n", ". ", " "]
    return _recursive_split(text, separators, chunk_size, overlap)


def _recursive_split(
    text: str,
    separators: list[str],
    chunk_size: int,
    overlap: int,
) -> list[str]:
    if not text.strip():
        return []

    if len(text) <= chunk_size:
        return [text.strip()]

    sep = separators[0] if separators else " "
    remaining_seps = separators[1:] if len(separators) > 1 else []

    parts = text.split(sep)
    chunks: list[str] = []
    current = ""

    for part in parts:
        candidate = f"{current}{sep}{part}" if current else part
        if len(candidate) <= chunk_size:
            current = candidate
        else:
            if current.strip():
                # If current chunk is too big, split further
                if len(current) > chunk_size and remaining_seps:
                    chunks.extend(
                        _recursive_split(current, remaining_seps, chunk_size, overlap)
                    )
                else:
                    chunks.append(current.strip())
            current = part

    if current.strip():
        if len(current) > chunk_size and remaining_seps:
            chunks.extend(
                _recursive_split(current, remaining_seps, chunk_size, overlap)
            )
        else:
            chunks.append(current.strip())

    # Add overlap between chunks
    if overlap > 0 and len(chunks) > 1:
        overlapped: list[str] = [chunks[0]]
        for i in range(1, len(chunks)):
            prev_tail = chunks[i - 1][-overlap:]
            overlapped.append(prev_tail + " " + chunks[i])
        chunks = overlapped

    return chunks


# ---------------------------------------------------------------------------
# Stage 4+5 — Embed & Store
# ---------------------------------------------------------------------------

async def index_document(
    db: AsyncSession,
    filepath: str,
    source_type: str = "official_docs",
    category: str = "general",
    url: str | None = None,
    chunk_size: int = 1200,
    overlap: int = 100,
) -> tuple[int, int]:
    """Full indexing pipeline: load → clean → chunk → embed → store.

    Uses async file loading (thread pool) and content hashing for
    incremental re-indexing.  Returns (article_id, chunks_created).
    """
    # Compute content hash for change detection
    file_hash = await asyncio.get_running_loop().run_in_executor(
        _file_parse_pool, compute_content_hash, filepath,
    )

    # Check if already indexed with same content
    existing = await _find_article_by_hash(db, file_hash)
    if existing is not None:
        logger.debug(
            "File %s unchanged (hash=%s…), skipping re-index",
            filepath, file_hash[:12],
        )
        return (existing.id, 0)

    # Load & clean (in thread pool to avoid blocking event loop)
    raw_text = await load_document_async(filepath)
    cleaned = clean_text(raw_text)

    if not cleaned:
        logger.warning("No text extracted from %s", filepath)
        return (0, 0)

    # Chunk
    chunks = chunk_text(cleaned, chunk_size=chunk_size, overlap=overlap)
    if not chunks:
        return (0, 0)

    logger.info("Indexing %s: %d chunks from %s", source_type, len(chunks), filepath)

    # Create parent article
    article = KnowledgeBaseArticle(
        category=category,
        question=Path(filepath).stem,
        answer=cleaned[:500],
        source_type=source_type,
        priority=10 if source_type == "official_docs" else 1,
        file_path=filepath,
        url=url,
        content_hash=file_hash,
    )
    db.add(article)
    await db.flush()

    # Embed all chunks
    embeddings = await embed_batch(chunks)

    # Compute expires_at for support_history
    expires_at = None
    if source_type == "support_history":
        expires_at = datetime.now(timezone.utc) + timedelta(
            days=settings.kb_retention_days
        )

    priority = 10 if source_type == "official_docs" else 1

    # Store chunks
    for idx, (chunk_text_content, emb) in enumerate(zip(chunks, embeddings)):
        chunk = KBChunk(
            article_id=article.id,
            content=chunk_text_content,
            chunk_index=idx,
            embedding=emb,
            source_type=source_type,
            priority=priority,
            expires_at=expires_at,
            metadata_={"filename": Path(filepath).name},
        )
        db.add(chunk)

    await db.flush()
    logger.info(
        "Indexed %d chunks for article_id=%d from %s (hash=%s…)",
        len(chunks), article.id, filepath, file_hash[:12],
    )
    return (article.id, len(chunks))


async def index_text(
    db: AsyncSession,
    text: str,
    source_type: str = "support_history",
    category: str = "general",
    priority: int = 1,
    title: str = "support_ticket",
    chunk_size: int = 1200,
    overlap: int = 100,
) -> tuple[int, int]:
    """Index raw text directly (e.g. from a closed ticket's Q+A).

    Uses content hashing to skip re-indexing identical text.
    """
    cleaned = clean_text(text)
    if not cleaned:
        return (0, 0)

    # Content hash for dedup
    text_hash = compute_text_hash(cleaned)
    existing = await _find_article_by_hash(db, text_hash)
    if existing is not None:
        logger.debug("Text already indexed (hash=%s…), skipping", text_hash[:12])
        return (existing.id, 0)

    chunks = chunk_text(cleaned, chunk_size=chunk_size, overlap=overlap)
    if not chunks:
        return (0, 0)

    article = KnowledgeBaseArticle(
        category=category,
        question=title,
        answer=cleaned[:500],
        source_type=source_type,
        priority=priority,
        content_hash=text_hash,
    )
    db.add(article)
    await db.flush()

    embeddings = await embed_batch(chunks)

    expires_at = None
    if source_type == "support_history":
        expires_at = datetime.now(timezone.utc) + timedelta(
            days=settings.kb_retention_days
        )

    for idx, (chunk_content, emb) in enumerate(zip(chunks, embeddings)):
        chunk = KBChunk(
            article_id=article.id,
            content=chunk_content,
            chunk_index=idx,
            embedding=emb,
            source_type=source_type,
            priority=priority,
            expires_at=expires_at,
        )
        db.add(chunk)

    await db.flush()
    return (article.id, len(chunks))


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

async def _find_article_by_hash(
    db: AsyncSession,
    content_hash: str,
) -> KnowledgeBaseArticle | None:
    """Find an existing article by its content hash (for incremental re-index)."""
    from sqlalchemy import select

    result = await db.execute(
        select(KnowledgeBaseArticle)
        .where(KnowledgeBaseArticle.content_hash == content_hash)
        .limit(1)
    )
    return result.scalar_one_or_none()
